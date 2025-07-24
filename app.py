import streamlit as st
import re
import joblib
import os
import pandas as pd
import matplotlib.pyplot as plt
from langchain.chat_models import ChatOpenAI
from langchain.prompts import PromptTemplate
from io import BytesIO
from docx import Document
from fpdf import FPDF
from PIL import Image
import numpy as np
import PyPDF2
import seaborn as sns
from sklearn.metrics import accuracy_score, confusion_matrix, classification_report


# Set page config

st.set_page_config(page_title="Smart Document Analysis", layout="wide")



def clean_text(text):
    text = text.lower()
    text = re.sub(r"[^a-zA-Z0-9\s]","",text)

    return text


def extract_text(uploaded_file):
    text = ""

    try:
        if uploaded_file.name.endswith(".txt"):
            text = uploaded_file.read().decode("utf-8")

        elif uploaded_file.name.endswith(".docx"):
            doc = Document(uploaded_file)
            text = "\n".join([para.text for para in doc.paragraphs])

        elif uploaded_file.name.endswith("pdf"):
            reader = PyPDF2.PdfReader(uploaded_file)
            text = "\n".join([page.extract_text() for page in reader.pages])
        
        else:
            st.error("Unsupported file type! Please upload .txt,.docx, or .pdf")
            return None
    except Exception as e:
        st.error(f"Error reading file: {str(e)}")
        return None
    
    return text


@st.cache_resource
def load_assets():
    
    return {
        'vectorizer':joblib.load("notebooks/tfidf_vectorizer.joblib"),
        'models' : {
            "Logistic Regression":joblib.load("models/logistic_regression.pkl"),
            "SVM": joblib.load("models/svm.pkl"),
            "Random Forest": joblib.load("models/random_forest.pkl")
        }

    }
    

def generate_llm_analysis(text,category,tempeature = 0.7):
    llm = ChatOpenAI(
        model= "meta-llama/Llama-3-8b-chat-hf",
        temperature = tempeature,
        openai_api_key = st.secrets["TOGETHER_API_KEY"],
        openai_api_base = "https://api.together.xyz/v1"
    )

    prompt = PromptTemplate.from_template("""
    Document Analysis:
    {text}
    
    Predicted Category: {category}
    
    1. Generate a suitable title
    2. Explain why this title fits
    3. Explain why the {category} category is appropriate
    """)
    
    response = llm.invoke(prompt.format(
        text=text[:2000],  # Limit text length
        category=category
    ))
    
    return response.content

def create_pdf_report(predictions, llm_analysis = None):
    pdf = FPDF ()
    pdf.add_page()
    pdf.set_font("Arial", size =14)


    # Title 
    pdf.cell(200, 10, txt="Document Classification Report", ln=1, align='C')
    pdf.ln(10)

    #Model predcitions

    pdf.set_font("Arial", size =12, style="B")
    pdf.cell(200, 10, txt="Model Predictions:", ln=1)
    pdf.set_font("Arial", size=10)

    for pred in predictions:
        pdf.cell(200, 10, 
                txt=f"{pred['Model']}: {pred['Prediction']} ({pred['Confidence']})", 
                ln=1)
    
    # LLM Analysis
    if llm_analysis:
        pdf.ln(10)
        pdf.set_font("Arial", size=12, style='B')
        pdf.cell(200, 10, txt="LLM Analysis:", ln=1)
        pdf.set_font("Arial", size=10)
        pdf.multi_cell(0, 10, txt=llm_analysis)
    
    return pdf.output(dest='S').encode('latin1')


def main():
    st.title("Smart Document Classifier with LLM Analysis")

    uploaded_file = st.file_uploader("Upload document", type=['txt','docx','pdf'])

    if uploaded_file:
        raw_text = extract_text(uploaded_file)
        cleaned_text = clean_text(raw_text)


        # Show Preview of the text
        with st.expander("📄 Document Preview"):
            st.text(raw_text[:500]+ "..."  if len(raw_text)>500 else raw_text)


        # Load models
        assets = load_assets()
        text_vec = assets['vectorizer'].transform([cleaned_text])


        # Make Category Prediction
        st.subheader("Model Prediction")
        result= []

        for name, model in assets["models"].items():
            prediction = model.predict(text_vec)[0]
            prob = max(model.predict_proba(text_vec)[0])
            result.append(
                {
                    "Model":name,
                    "Prediction":prediction,
                    "Confidence":f"{prob:.2%}"
                }
            )

        result_df = pd.DataFrame(result)
        st.dataframe(result_df)

        best_result = max(result, key=lambda x : x["Confidence"])
        st.markdown(f"### 🤖 Best Prediction: **{best_result["Prediction"]}** by {best_result['Model']} ({best_result['Confidence']}% confidence)")

        st.subheader("🤖 LLM Analysis")

        temp = st.slider("LLM Creativity (temperature)", min_value=0.1, max_value=1.0, value=0.7, step=0.1)


        # Select model manually by user
        selected_model = st.selectbox("Use prediction from: ",
        [res["Model"] for res in result])

        selected_pred = next(res["Prediction"] for res in result if res["Model"]==selected_model)

        if st.button("Generate Analysis"):
            with st.spinner("Generating LLM Analysis..."):
                try:
                    
                    response = generate_llm_analysis(raw_text, selected_pred, temp)
                    title = response.split("\n")[0]
                    analysis = "\n".join(response.split("\n")[1:])

                    st.write("### Generated Title:")
                    st.write(title)
                    st.write("###Analysis")
                    st.write(analysis)
                    pdf_report = create_pdf_report(result,llm_analysis=analysis)
                    st.download_button(
                    label="📥 Download PDF Report",
                    data=pdf_report,
                    file_name="classification_report.pdf",
                    mime="application/pdf"
                )
                except Exception as e:
                    st.error(f"LLM analysis failed: {str(e)}")


        
 

if __name__ == "__main__":
    main()